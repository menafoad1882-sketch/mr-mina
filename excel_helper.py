"""
مساعد توليد ملفات التقارير.
يستخدم openpyxl لتوليد Excel (.xlsx)، ولو المكتبة غير متوفرة في بيئة النشر
يرجع تلقائيًا لملف CSV بدلًا من أن يتعطّل الموقع (production-safe).
"""
import io
import csv

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_OPENPYXL = True
except Exception:  # pragma: no cover
    HAS_OPENPYXL = False


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv; charset=utf-8"


def build_report_file(title, headers, rows, footer=None, rtl=True):
    """
    يبني ملف تقرير ويعيد (bytes, filename_ext, mimetype).
    - title: عنوان التقرير (يظهر كصف أول في Excel)
    - headers: قائمة عناوين الأعمدة
    - rows: قائمة صفوف (كل صف قائمة قيم)
    - footer: صف إجمالي اختياري (قائمة قيم)
    يرجّع xlsx لو openpyxl متاح، وإلا csv.
    """
    if HAS_OPENPYXL:
        return _build_xlsx(title, headers, rows, footer, rtl)
    return _build_csv(title, headers, rows, footer)


def _build_xlsx(title, headers, rows, footer, rtl):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "التقرير"
    if rtl:
        ws.sheet_view.rightToLeft = True
    head_font = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2563EB")

    ws.append([title])
    ws.append(headers)
    for cell in ws[2]:
        cell.font = head_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(list(r))
    if footer is not None:
        ws.append([])
        ws.append(list(footer))
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    # عرض الأعمدة
    for i in range(len(headers)):
        col = openpyxl.utils.get_column_letter(i + 1)
        ws.column_dimensions[col].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue(), "xlsx", XLSX_MIME


def _build_csv(title, headers, rows, footer):
    sio = io.StringIO()
    sio.write("\ufeff")  # BOM عشان العربي يظهر صح في Excel
    w = csv.writer(sio)
    w.writerow([title])
    w.writerow(headers)
    for r in rows:
        w.writerow(list(r))
    if footer is not None:
        w.writerow([])
        w.writerow(list(footer))
    data = sio.getvalue().encode("utf-8")
    return data, "csv", CSV_MIME
